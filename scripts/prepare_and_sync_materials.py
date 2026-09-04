# -*- coding: utf-8 -*-
"""
ベストワン教材一覧 Excel を CataBooks の database.js 形式へ変換し、
必要なら Supabase の materials へ upsert する。

iワーク（i-ワーク / ｉワーク 含む）は Excel に無くても旧マスタから残す。
カート・お気に入りで参照中の教材は削除しない。
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.parse import quote
from urllib.request import Request, urlopen

try:
    from openpyxl import load_workbook
except ImportError:
    print("openpyxl が必要です: pip install openpyxl")
    sys.exit(1)

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = Path(
    r"C:\Users\user\OneDrive\事務書類\_事務資料\_テキスト請求業務\2026年度\ベストワン教材一覧_20260904.xlsx"
)
DB_JS = ROOT / "database.js"
CLIENT_JS = ROOT / "supabase-client.js"


def to_int(value: object) -> int:
    if value is None or value == "":
        return 0
    if isinstance(value, bool):
        return 0
    if isinstance(value, (int, float)):
        return int(round(float(value)))
    text = str(value).replace("¥", "").replace(",", "").replace("円", "").strip()
    if text in ("", "-", "—"):
        return 0
    try:
        return int(round(float(text)))
    except ValueError:
        return 0


def header_index(headers: list[object], *names: str) -> int | None:
    normalized = {str(h).strip(): i + 1 for i, h in enumerate(headers) if h}
    for name in names:
        if name in normalized:
            return normalized[name]
    return None


def is_iwork(title: str) -> bool:
    # 長音「ー」は「ワーク」の一部なので消さない。空白とハイフンだけ除く。
    compact = unicodedata.normalize("NFKC", title or "").lower()
    compact = re.sub(r"[\s\u3000\-‐‑–—_]", "", compact)
    return "iワーク" in compact or "ｉワーク" in compact


def load_old_materials(path: Path) -> list[dict]:
    if not path.exists():
        return []
    text = path.read_text(encoding="utf-8")
    match = re.search(r"=\s*(\[.*\])\s*;?\s*$", text, re.S)
    if not match:
        raise RuntimeError(f"{path} から配列を読み取れませんでした")
    return json.loads(match.group(1))


def read_excel(src: Path) -> list[dict]:
    wb = load_workbook(src, data_only=True)
    sheet_name = "教材一覧" if "教材一覧" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]

    col_id = header_index(headers, "物品コード", "id", "教材ID", "教材コード")
    col_title = header_index(headers, "物品名", "title", "テキスト名", "教材名")
    col_wholesale = header_index(headers, "提供価格", "提供価格数値", "price_wholesale", "仕入価格")
    col_retail = header_index(headers, "生徒販売価格", "生徒販売価格数値", "price_retail", "販売価格")

    if col_title is None:
        col_id, col_title, col_wholesale, col_retail = 1, 2, 3, 4
    else:
        col_id = col_id or 2
        col_wholesale = col_wholesale or 8
        col_retail = col_retail or 9

    rows: list[dict] = []
    seen: set[str] = set()
    for r in range(2, ws.max_row + 1):
        title = ws.cell(r, col_title).value
        if not title or not str(title).strip():
            continue
        mid = ws.cell(r, col_id).value if col_id else None
        item_id = str(mid).strip() if mid else f"NOID-{r}"
        if item_id in seen:
            continue
        seen.add(item_id)
        rows.append(
            {
                "title": str(title).strip(),
                "price_retail": to_int(ws.cell(r, col_retail).value) if col_retail else 0,
                "id": item_id,
                "price_wholesale": to_int(ws.cell(r, col_wholesale).value) if col_wholesale else 0,
            }
        )
    return rows


def merge_keep_iwork(excel_rows: list[dict], old_rows: list[dict]) -> tuple[list[dict], dict]:
    by_id = {row["id"]: row for row in excel_rows}
    excel_iwork = [row for row in excel_rows if is_iwork(row["title"])]
    old_iwork = [row for row in old_rows if is_iwork(row.get("title", ""))]
    kept_from_old = []
    for row in old_iwork:
        if row["id"] in by_id:
            continue
        kept_from_old.append(
            {
                "title": row["title"],
                "price_retail": int(row.get("price_retail") or 0),
                "id": row["id"],
                "price_wholesale": int(row.get("price_wholesale") or 0),
            }
        )
        by_id[row["id"]] = kept_from_old[-1]

    merged = list(by_id.values())
    stats = {
        "excel": len(excel_rows),
        "excel_iwork": len(excel_iwork),
        "old": len(old_rows),
        "old_iwork": len(old_iwork),
        "kept_iwork": len(kept_from_old),
        "merged": len(merged),
        "kept_iwork_titles": [r["title"] for r in kept_from_old],
        "excel_iwork_titles": [r["title"] for r in excel_iwork],
        "old_iwork_titles": [r["title"] for r in old_iwork],
    }
    return merged, stats


def write_database_js(rows: list[dict], path: Path) -> None:
    payload = json.dumps(rows, ensure_ascii=False, indent=4)
    path.write_text(f"const DB_DATA = {payload};\n", encoding="utf-8")


def read_supabase_creds() -> tuple[str, str]:
    text = CLIENT_JS.read_text(encoding="utf-8")
    url = re.search(r"const SUPABASE_URL = '([^']+)'", text)
    key = re.search(r"const SUPABASE_ANON_KEY = '([^']+)'", text)
    if not url or not key:
        raise RuntimeError("supabase-client.js から接続情報を読めませんでした")
    return url.group(1), key.group(1)


def rest(url: str, key: str, method: str, path: str, body: object | None = None, extra: dict | None = None):
    headers = {
        "apikey": key,
        "Authorization": f"Bearer {key}",
        "Content-Type": "application/json",
        "Prefer": "return=representation",
    }
    if extra:
        headers.update(extra)
    data = None if body is None else json.dumps(body).encode("utf-8")
    req = Request(url.rstrip("/") + path, data=data, headers=headers, method=method)
    try:
        with urlopen(req, timeout=120) as resp:
            raw = resp.read().decode("utf-8")
            return json.loads(raw) if raw else None
    except HTTPError as e:
        detail = e.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"HTTP {e.code} {path}: {detail[:500]}") from e
    except URLError as e:
        raise RuntimeError(f"接続失敗 {path}: {e}") from e


def fetch_all(url: str, key: str, table: str, select: str) -> list[dict]:
    rows: list[dict] = []
    page = 1000
    start = 0
    while True:
        chunk = rest(
            url,
            key,
            "GET",
            f"/rest/v1/{table}?select={select}&order=id",
            extra={"Range": f"{start}-{start + page - 1}", "Prefer": "count=exact"},
        )
        if not chunk:
            break
        rows.extend(chunk)
        if len(chunk) < page:
            break
        start += page
    return rows


def upsert_materials(url: str, key: str, rows: list[dict]) -> int:
    done = 0
    chunk_size = 400
    for i in range(0, len(rows), chunk_size):
        chunk = rows[i : i + chunk_size]
        rest(
            url,
            key,
            "POST",
            "/rest/v1/materials",
            chunk,
            extra={"Prefer": "resolution=merge-duplicates,return=minimal"},
        )
        done += len(chunk)
        print(f"  upsert {done}/{len(rows)}")
    return done


def delete_ids(url: str, key: str, ids: list[str]) -> int:
    deleted = 0
    chunk_size = 40
    for i in range(0, len(ids), chunk_size):
        chunk = ids[i : i + chunk_size]
        encoded = ",".join(quote(item, safe="") for item in chunk)
        rest(
            url,
            key,
            "DELETE",
            f"/rest/v1/materials?id=in.({encoded})",
            extra={"Prefer": "return=minimal"},
        )
        deleted += len(chunk)
        print(f"  delete {deleted}/{len(ids)}")
    return deleted


def print_stats(label: str, rows: list[dict]) -> None:
    iwork = [r for r in rows if is_iwork(r.get("title", ""))]
    print(f"{label}: {len(rows)} 件 / iワーク {len(iwork)} 件")
    for row in iwork:
        print(f"  - {row.get('id')} / {row.get('title')}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("xlsx", nargs="?", default=str(DEFAULT_XLSX))
    parser.add_argument("--sync", action="store_true", help="Supabase へ upsert（iワーク以外の欠落は prune）")
    args = parser.parse_args()

    src = Path(args.xlsx)
    if not src.exists():
        print(f"Excel が見つかりません: {src}")
        sys.exit(1)

    excel_rows = read_excel(src)
    old_rows = load_old_materials(DB_JS)
    merged, stats = merge_keep_iwork(excel_rows, old_rows)
    write_database_js(merged, DB_JS)

    print(f"Excel: {stats['excel']} 件（iワーク {stats['excel_iwork']} 件）")
    print(f"旧 database.js: {stats['old']} 件（iワーク {stats['old_iwork']} 件）")
    print(f"Excelに無く残した iワーク: {stats['kept_iwork']} 件")
    for title in stats["kept_iwork_titles"]:
        print(f"  keep: {title}")
    for title in stats["excel_iwork_titles"]:
        print(f"  excel iワーク: {title}")
    print(f"書き出し後 database.js: {stats['merged']} 件 -> {DB_JS}")

    if not args.sync:
        return

    url, key = read_supabase_creds()
    current = fetch_all(url, key, "materials", "id,title,price_wholesale,price_retail")
    carts = fetch_all(url, key, "cart_items", "material_id")
    favs = fetch_all(url, key, "favorites", "material_id")
    print_stats("本番DB（更新前）", current)

    current_iwork = [r for r in current if is_iwork(r.get("title", ""))]
    extra_keep = []
    for row in current_iwork:
        if row["id"] not in {m["id"] for m in merged}:
            extra = {
                "title": row["title"],
                "price_retail": int(row.get("price_retail") or 0),
                "id": row["id"],
                "price_wholesale": int(row.get("price_wholesale") or 0),
            }
            merged.append(extra)
            extra_keep.append(extra)
    if extra_keep:
        write_database_js(merged, DB_JS)
        print(f"本番DBから追加で残した iワーク: {len(extra_keep)} 件")

    referenced = {r["material_id"] for r in carts if r.get("material_id")}
    referenced |= {r["material_id"] for r in favs if r.get("material_id")}
    keep_ids = {m["id"] for m in merged} | referenced
    prune_ids = [r["id"] for r in current if r["id"] not in keep_ids]
    prune_iwork = [r for r in current if r["id"] in prune_ids and is_iwork(r.get("title", ""))]
    if prune_iwork:
        raise RuntimeError("iワークを prune 対象にしてしまったため中止します")

    print(f"upsert 対象: {len(merged)} 件")
    upsert_materials(url, key, merged)
    print(f"カート/お気に入り参照のため残す ID: {len(referenced)} 件")
    print(f"prune 対象（iワーク以外・参照なし）: {len(prune_ids)} 件")
    if prune_ids:
        delete_ids(url, key, prune_ids)

    after = fetch_all(url, key, "materials", "id,title,price_wholesale,price_retail")
    print_stats("本番DB（更新後）", after)
    target = next((r for r in after if "標準新演習" in r.get("title", "") and "中３" in r.get("title", "") and "英語" in r.get("title", "") and "ＸY" in r.get("title", "")), None)
    if target:
        print(f"確認: {target['id']} {target['title']} 販売{target.get('price_retail')} 仕入{target.get('price_wholesale')}")
    else:
        print("確認: 標準新演習　中３　英語　ＸY が見つかりません")


if __name__ == "__main__":
    main()
